import io
import unittest

import pandas as pd

from exporter import export_to_csv, export_to_excel, export_to_word


class ExporterTests(unittest.TestCase):
    def test_export_data_kosong_tanpa_kolom(self):
        hasil = export_to_excel(pd.DataFrame())
        self.assertGreater(hasil.getbuffer().nbytes, 0)

    def test_semua_format_export(self):
        data = pd.DataFrame([
            {
                "tanggal": "2026-10-24",
                "nama_menu": "Menu Uji",
                "energi_kcal": 100.0,
                "protein_g": 10.0,
            }
        ])
        self.assertTrue(export_to_csv(data).startswith(b"tanggal"))
        self.assertGreater(export_to_excel(data).getbuffer().nbytes, 0)
        self.assertGreater(
            export_to_word(data, "2026-10-24", "2026-10-24").getbuffer().nbytes,
            0,
        )

    def test_excel_dengan_lembar_akg(self):
        from openpyxl import load_workbook

        data = pd.DataFrame([
            {
                "tanggal": "2026-10-24",
                "nama_menu": "Menu Uji",
                "energi_kcal": 100.0,
                "protein_g": 10.0,
            }
        ])
        df_akg = pd.DataFrame([
            {
                "tanggal": "2026-10-24",
                "energi_pct": 50.0,
                "protein_pct": 25.0,
                "lemak_pct": 80.0,
                "karbo_pct": 40.0,
                "penilaian": "Sangat kurang",
            }
        ])
        hasil = export_to_excel(data, akg=("7-9 tahun", df_akg))
        wb = load_workbook(io.BytesIO(hasil.getvalue()))
        self.assertIn("Kecukupan AKG", wb.sheetnames)
        self.assertIn("Rekap Harian", wb.sheetnames)
        # Tanpa argumen akg, lembar AKG tidak dibuat (kompatibel mundur).
        hasil2 = export_to_excel(data)
        wb2 = load_workbook(io.BytesIO(hasil2.getvalue()))
        self.assertNotIn("Kecukupan AKG", wb2.sheetnames)


if __name__ == "__main__":
    unittest.main()
